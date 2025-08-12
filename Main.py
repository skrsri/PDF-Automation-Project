import os
import pandas as pd
import tkinter as tk
from tkinter import ttk
from tkinter import messagebox
from docx import Document
from docx2pdf import convert

# Set paths
EXCEL_PATH = os.path.abspath("template.xlsx")
TEMPLATE_2026 = os.path.abspath("2026_Invitation_MMMUT.docx")
TEMPLATE_BPHARMA = os.path.abspath("Bpharma_Invitation_MMMUT.docx")
OUTPUT_DIR = os.path.abspath("output_pdfs")
os.makedirs(OUTPUT_DIR, exist_ok=True)

# Read Excel
data = pd.read_excel(EXCEL_PATH)
data.columns = data.columns.str.strip().str.lower()

# Replace placeholders
def replace_placeholders(doc, replacements):
    for para in doc.paragraphs:
        for placeholder, replacement in replacements.items():
            if placeholder in para.text:
                for run in para.runs:
                    if placeholder in run.text:
                        run.text = run.text.replace(placeholder, replacement)
                        run.bold = True

# Process each row
def process_invitations(template_path):
    for index, row in data.iterrows():
        company_name = row["company name"].strip()
        company_folder = os.path.join(OUTPUT_DIR, company_name)
        os.makedirs(company_folder, exist_ok=True)

        doc = Document(template_path)

        name = row["hr name"].strip().lower()
        if "mr" in name:
            salutation = "Dear Sir,"
        elif "ms" in name:
            salutation = "Dear Ma'am,"
        else:
            salutation = "Dear Sir/Ma'am"

        replacements = {
            "{HR_NAME}": row["hr name"],
            "{DESIGNATION}": row["designation"],
            "{COMPANY_NAME}": row["company name"],
            "{GENDER}": row["gender"],
            "{s}": salutation,
        }

        replace_placeholders(doc, replacements)

        subfolder = os.path.join(company_folder, f"{index+1}")
        os.makedirs(subfolder, exist_ok=True)

        docx_path = os.path.join(subfolder, "temp.docx")
        pdf_path = os.path.join(subfolder, f"{company_name}_Invitation_MMMUT.pdf")

        doc.save(docx_path)
        convert(docx_path, pdf_path)
        os.remove(docx_path)

    messagebox.showinfo("Success", "Invitations generated successfully.")

# UI setup
def run_ui():
    root = tk.Tk()
    root.title("Invitation Generator")
    root.configure(bg="#121212")
    root.geometry("400x250")

    style = ttk.Style(root)
    style.theme_use("clam")
    style.configure("TRadiobutton", foreground="white", background="#121212", font=("Segoe UI", 12))
    style.configure("TButton", foreground="black", font=("Segoe UI", 12), padding=6)

    selected_template = tk.StringVar(value="2026")

    label = ttk.Label(root, text="Select Invitation Template", background="#121212", foreground="white", font=("Segoe UI", 14))
    label.pack(pady=20)

    ttk.Radiobutton(root, text="2026 Batch", variable=selected_template, value="2026").pack()
    ttk.Radiobutton(root, text="Bpharma", variable=selected_template, value="Bpharma").pack()

    def generate():
        if selected_template.get() == "2026":
            template = TEMPLATE_2026
        else:
            template = TEMPLATE_BPHARMA

        if not os.path.exists(template):
            messagebox.showerror("Error", "Template file not found.")
            return

        process_invitations(template)

    ttk.Button(root, text="Generate Invitations", command=generate).pack(pady=30)
    root.mainloop()

if __name__ == "__main__":
    run_ui()
